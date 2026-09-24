#!/usr/bin/env python3
"""Conversor Markdown -> .docx para business/documentacao/.

Fonte unica sancionada dos .docx derivados. O .md e a fonte de verdade; este
script e a unica forma legitima de produzir o binario correspondente. Nao editar
os .docx a mao, nem por OOXML.

Uso
---
    python business/scripts/gerar-docx.py              # regenera os seis
    python business/scripts/gerar-docx.py --verificar  # so verifica, nao escreve
    python business/scripts/gerar-docx.py roadmap-tecnologico.md   # so um

Dependencias
------------
    Python >= 3.10   (validado em 3.14.5)
    python-docx == 1.2.0
Nao usa pandoc, nem Word, nem LibreOffice.

Familia tipografica
-------------------
"LexCV Institucional" — a familia originalmente usada em termo-de-abertura.docx,
escolhida por ser a unica das tres em uso que define de facto um sistema (corpo
Calibri 11/1.15, rampa de titulos em azul-marinho, regua horizontal fina,
caixa de citacao com barra lateral, tabela com cabecalho sombreado). As outras
duas eram o template de origem do python-docx sem qualquer ajuste, e uma variante
A4-paisagem a 10pt. Ver ESPECIFICACAO no fim deste ficheiro.

Formato de pagina
-----------------
A4 (210x297 mm) nos seis. Os documentos destinam-se a clientes institucionais em
Cabo Verde, onde o tabuleiro corrente e A4; US Letter sairia com margens erradas
na impressao. A familia define tipos, cores, estilos e espacamentos — nao define
orientacao. ORIENTACAO_PAISAGEM lista os documentos cuja grelha de tabelas nao
cabe em retrato; hoje so o roadmap-tecnologico. Tudo o resto (fontes, cores,
rampa de titulos, margens, corpo, entrelinha) e identico nos seis.

Suporte de Markdown
-------------------
Titulos 1-4, listas em marcas (2 niveis) e numeradas, tabelas com cabecalho,
citacoes (incluindo listas dentro da citacao), reguas horizontais, **negrito**,
*italico*, ***ambos***, `codigo` e [ligacoes](url). Nao ha blocos de codigo
cercados, imagens nem HTML nos seis documentos-fonte; nao sao suportados.

Invariante de espacamento
-------------------------
Nenhum paragrafo produzido pode conter uma sequencia de dois ou mais espacos.
Foi esse o defeito de uma geracao anterior por edicao directa de OOXML. O modo
--verificar falha em erro se voltar a acontecer.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

try:
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Emu, Inches, Mm, Pt, RGBColor, Twips
except ImportError:  # pragma: no cover
    sys.exit(
        "python-docx nao esta instalado.\n"
        "  pip install --user 'python-docx==1.2.0'"
    )

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


# --------------------------------------------------------------------------
# ESPECIFICACAO da familia "LexCV Institucional"
# --------------------------------------------------------------------------

FONTE_TEXTO = "Calibri"
FONTE_MONO = "Consolas"

AZUL_TITULO_1 = RGBColor(0x0F, 0x27, 0x44)
AZUL_TITULO_N = RGBColor(0x1B, 0x4D, 0x8F)
CINZA_REGUA = "B8C4D4"
AZUL_BARRA_CITACAO = "1B4D8F"
FUNDO_CABECALHO_TABELA = "E8EEF6"

A4_CURTA = Mm(210)
A4_LONGA = Mm(297)
MARGEM_LATERAL = Inches(0.95)
MARGEM_VERTICAL = Inches(0.87)

# Documentos que saem em A4 paisagem. Nao e uma segunda familia tipografica: e
# a mesma familia numa folha rodada. O roadmap tem uma grelha de 6 colunas cuja
# coluna "Estado" ficava a 0,40 pol em retrato — cerca de cinco caracteres, com
# "Previsto" a nao caber. Ver ESPECIFICACAO no fim do ficheiro.
ORIENTACAO_PAISAGEM = frozenset({"roadmap-tecnologico.md"})

CORPO_PT = 11.0
CORPO_ENTRELINHA = 1.15
TABELA_PT = 10.0

# Largura de coluna: piso absoluto, e piso adicional para que a palavra mais
# longa da coluna caiba sem ser partida. As duas constantes em "em" sao uma
# estimativa deliberadamente generosa da largura de um caractere em Calibri,
# em fraccao do corpo. Conferidas contra as metricas reais da fonte:
#   "estimativa"        4,23 em real  ->  4,80 em estimado
#   "Previsto"          3,30 em real  ->  3,96 em estimado
#   "PLATAFORMA_ADMIN"  9,14 em real  ->  9,60 em estimado
# A estimativa erra sempre por excesso, nos tres regimes (minusculas, capitalizado,
# tudo em maiusculas). Nao ha dependencia de fonte instalada: o gerador continua
# a correr em qualquer maquina, sem ler ficheiros de tipos.
LARGURA_EM_MAIUSCULA = 0.60      # maiusculas, digitos e pontuacao
LARGURA_EM_MINUSCULA = 0.48
MARGENS_CELULA_TWIPS = 230       # 0,08 pol de cada lado, omissao do Word
PISO_COLUNA = 0.06               # fraccao da largura util
TETO_MINIMO_COLUNA = 0.25        # nenhum piso por palavra passa disto

# recuos, em twips (1/20 pt)
RECUO_CITACAO = 340
RECUO_LISTA_CITACAO = 620
RECUO_LISTA_N2 = 720

RAMPA_TITULOS = {
    # nivel: (tamanho pt, cor, espaco antes pt, espaco depois pt)
    1: (20.0, AZUL_TITULO_1, 0, 10),
    2: (15.0, AZUL_TITULO_N, 14, 6),
    3: (12.5, AZUL_TITULO_N, 12, 4),
    4: (11.5, AZUL_TITULO_N, 10, 3),
}


# --------------------------------------------------------------------------
# Analise do Markdown
# --------------------------------------------------------------------------

RE_TITULO = re.compile(r"^(#{1,6})\s+(.*)$")
RE_REGUA = re.compile(r"^\s{0,3}(?:-{3,}|\*{3,}|_{3,})\s*$")
RE_MARCA = re.compile(r"^(\s*)[-*+]\s+(.*)$")
RE_NUMERO = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
RE_SEPARADOR_TABELA = re.compile(r"^\s*\|?[\s:|-]*-[\s:|-]*\|?\s*$")


@dataclass
class Bloco:
    tipo: str  # titulo | paragrafo | marca | numero | citacao | tabela | regua
    texto: str = ""
    nivel: int = 0
    linhas: list[str] = field(default_factory=list)
    filhos: list["Bloco"] = field(default_factory=list)


def _e_linha_de_tabela(linhas: list[str], i: int) -> bool:
    if not linhas[i].lstrip().startswith("|"):
        return False
    if i + 1 >= len(linhas):
        return False
    seguinte = linhas[i + 1]
    return seguinte.lstrip().startswith("|") and bool(RE_SEPARADOR_TABELA.match(seguinte))


def _celulas(linha: str) -> list[str]:
    bruto = linha.strip()
    if bruto.startswith("|"):
        bruto = bruto[1:]
    if bruto.endswith("|"):
        bruto = bruto[:-1]
    # separador de celula = "|" nao escapado
    partes = re.split(r"(?<!\\)\|", bruto)
    return [p.strip().replace("\\|", "|") for p in partes]


def analisar(texto: str) -> list[Bloco]:
    """Markdown -> lista plana de blocos."""
    linhas = texto.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocos: list[Bloco] = []
    i = 0
    n = len(linhas)

    while i < n:
        linha = linhas[i]

        if not linha.strip():
            i += 1
            continue

        if RE_REGUA.match(linha):
            blocos.append(Bloco("regua"))
            i += 1
            continue

        m = RE_TITULO.match(linha)
        if m:
            blocos.append(Bloco("titulo", texto=m.group(2).strip(), nivel=len(m.group(1))))
            i += 1
            continue

        if _e_linha_de_tabela(linhas, i):
            j = i
            corpo: list[str] = []
            while j < n and linhas[j].lstrip().startswith("|"):
                corpo.append(linhas[j])
                j += 1
            blocos.append(Bloco("tabela", linhas=corpo))
            i = j
            continue

        if linha.startswith(">"):
            j = i
            interior: list[str] = []
            while j < n and linhas[j].startswith(">"):
                sem_marca = linhas[j][1:]
                if sem_marca.startswith(" "):
                    sem_marca = sem_marca[1:]
                interior.append(sem_marca)
                j += 1
            blocos.append(Bloco("citacao", filhos=analisar("\n".join(interior))))
            i = j
            continue

        m = RE_MARCA.match(linha)
        if m:
            recuo = len(m.group(1).expandtabs(4))
            item, i = _juntar_item(linhas, i, m.group(2))
            blocos.append(Bloco("marca", texto=item, nivel=2 if recuo >= 2 else 1))
            continue

        m = RE_NUMERO.match(linha)
        if m:
            recuo = len(m.group(1).expandtabs(4))
            item, i = _juntar_item(linhas, i, m.group(3))
            blocos.append(Bloco("numero", texto=item, nivel=2 if recuo >= 2 else 1))
            continue

        # paragrafo: linhas consecutivas ate branco ou inicio de outro bloco
        partes = [linha.strip()]
        i += 1
        while i < n and linhas[i].strip() and not _inicia_bloco(linhas[i]):
            partes.append(linhas[i].strip())
            i += 1
        blocos.append(Bloco("paragrafo", texto=" ".join(partes)))

    return blocos


def _inicia_bloco(linha: str) -> bool:
    return bool(
        RE_TITULO.match(linha)
        or RE_REGUA.match(linha)
        or RE_MARCA.match(linha)
        or RE_NUMERO.match(linha)
        or linha.startswith(">")
        or linha.lstrip().startswith("|")
    )


def _juntar_item(linhas: list[str], i: int, primeiro: str) -> tuple[str, int]:
    """Junta a continuacao recuada de um item de lista numa so linha logica."""
    partes = [primeiro.strip()]
    i += 1
    n = len(linhas)
    while i < n:
        seguinte = linhas[i]
        if not seguinte.strip():
            break
        if RE_MARCA.match(seguinte) or RE_NUMERO.match(seguinte) or _inicia_bloco(seguinte):
            break
        if not seguinte.startswith((" ", "\t")):
            break
        partes.append(seguinte.strip())
        i += 1
    return " ".join(partes), i


# --------------------------------------------------------------------------
# Analise inline
# --------------------------------------------------------------------------

RE_INLINE = re.compile(
    r"(?P<codigo>`+[^`]+`+)"
    r"|(?P<ligacao>\[(?P<lig_texto>[^\]]+)\]\((?P<lig_url>[^)\s]+)\))"
    r"|(?P<bi>\*\*\*(?P<bi_t>[^*]+)\*\*\*)"
    r"|(?P<negrito>\*\*(?P<neg_t>(?:[^*]|\*(?!\*))+)\*\*)"
    r"|(?P<italico>(?<![\w*])\*(?P<ita_t>[^*\s](?:[^*]*[^*\s])?)\*(?![\w*]))"
    r"|(?P<italico_>(?<![\w_])_(?P<ita_t_>[^_\s](?:[^_]*[^_\s])?)_(?![\w_]))"
)


@dataclass
class Fragmento:
    texto: str
    negrito: bool = False
    italico: bool = False
    codigo: bool = False


def analisar_inline(texto: str) -> list[Fragmento]:
    frags: list[Fragmento] = []
    pos = 0
    for m in RE_INLINE.finditer(texto):
        if m.start() > pos:
            frags.append(Fragmento(texto[pos:m.start()]))
        if m.group("codigo"):
            frags.append(Fragmento(m.group("codigo").strip("`"), codigo=True))
        elif m.group("ligacao"):
            rotulo, url = m.group("lig_texto"), m.group("lig_url")
            frags.extend(analisar_inline(rotulo))
            if url not in rotulo:
                frags.append(Fragmento(f" ({url})"))
        elif m.group("bi"):
            frags.append(Fragmento(m.group("bi_t"), negrito=True, italico=True))
        elif m.group("negrito"):
            frags.extend(_herdar(analisar_inline(m.group("neg_t")), negrito=True))
        elif m.group("italico"):
            frags.extend(_herdar(analisar_inline(m.group("ita_t")), italico=True))
        elif m.group("italico_"):
            frags.extend(_herdar(analisar_inline(m.group("ita_t_")), italico=True))
        pos = m.end()
    if pos < len(texto):
        frags.append(Fragmento(texto[pos:]))
    return _normalizar_espacos(frags)


def _herdar(frags: list[Fragmento], *, negrito: bool = False, italico: bool = False):
    for f in frags:
        f.negrito = f.negrito or negrito
        f.italico = f.italico or italico
    return frags


def _normalizar_espacos(frags: list[Fragmento]) -> list[Fragmento]:
    """Garante zero sequencias de espacos multiplos. Codigo fica intacto.

    Invariante do modulo: nenhum paragrafo escrito no .docx contem "  ".
    """
    saida: list[Fragmento] = []
    for f in frags:
        if f.codigo:
            if f.texto:
                saida.append(f)
            continue
        f.texto = re.sub(r"[\s ]+", " ", f.texto)
        if f.texto:
            saida.append(f)
    # colar fronteiras: " " + " x" nao pode dar "  x"
    for a, b in zip(saida, saida[1:]):
        if a.texto.endswith(" ") and b.texto.startswith(" ") and not b.codigo:
            b.texto = b.texto.lstrip(" ")
    if saida:
        saida[0].texto = saida[0].texto.lstrip(" ")
        saida[-1].texto = saida[-1].texto.rstrip(" ")
    return [f for f in saida if f.texto]


# --------------------------------------------------------------------------
# Construcao do .docx
# --------------------------------------------------------------------------

def _borda(pPr, lado: str, **attrs):
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = pPr.makeelement(qn("w:pBdr"), {})
        pPr.append(pBdr)
    el = pBdr.makeelement(qn(f"w:{lado}"), {})
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    pBdr.append(el)


def _sombrear(celula, cor_hex: str):
    tcPr = celula._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    tcPr.append(shd)


def _repetir_cabecalho(linha):
    trPr = linha._tr.get_or_add_trPr()
    el = trPr.makeelement(qn("w:tblHeader"), {})
    trPr.append(el)


def dimensoes(paisagem: bool) -> tuple[int, int]:
    """(largura, altura) da folha A4 na orientacao pedida, em EMU."""
    return (A4_LONGA, A4_CURTA) if paisagem else (A4_CURTA, A4_LONGA)


def preparar_documento(paisagem: bool = False) -> "docx.document.Document":
    doc = docx.Document()

    largura, altura = dimensoes(paisagem)
    sec = doc.sections[0]
    # python-docx nao troca as dimensoes sozinho ao mudar `orientation`; as tres
    # linhas seguintes tem de andar juntas ou o Word imprime retrato na mesma.
    sec.orientation = WD_ORIENT.LANDSCAPE if paisagem else WD_ORIENT.PORTRAIT
    sec.page_width = largura
    sec.page_height = altura
    sec.left_margin = sec.right_margin = MARGEM_LATERAL
    sec.top_margin = sec.bottom_margin = MARGEM_VERTICAL

    normal = doc.styles["Normal"]
    normal.font.name = FONTE_TEXTO
    normal.font.size = Pt(CORPO_PT)
    normal.paragraph_format.line_spacing = CORPO_ENTRELINHA
    normal.paragraph_format.space_after = Pt(6)
    # garantir a fonte tambem para o conjunto East Asian / complexo
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONTE_TEXTO)

    for nivel, (tam, cor, antes, depois) in RAMPA_TITULOS.items():
        st = doc.styles[f"Heading {nivel}"]
        st.font.name = FONTE_TEXTO
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.color.rgb = cor
        st.paragraph_format.space_before = Pt(antes)
        st.paragraph_format.space_after = Pt(depois)
        st.paragraph_format.keep_with_next = True

    for nome, recuo in (("List Bullet", None), ("List Bullet 2", RECUO_LISTA_N2),
                        ("List Number", None), ("List Number 2", RECUO_LISTA_N2)):
        st = doc.styles[nome]
        st.font.name = FONTE_TEXTO
        st.font.size = Pt(CORPO_PT)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = CORPO_ENTRELINHA
        if recuo:
            st.paragraph_format.left_indent = Twips(recuo)

    return doc


def escrever_fragmentos(par, frags: list[Fragmento]):
    for f in frags:
        run = par.add_run(f.texto)
        run.bold = f.negrito or None
        run.italic = f.italico or None
        if f.codigo:
            run.font.name = FONTE_MONO
            run.font.size = Pt(CORPO_PT - 1)
            run.font.color.rgb = AZUL_TITULO_1
            run._element.rPr.rFonts.set(qn("w:cs"), FONTE_MONO)


def _largura_util(paisagem: bool = False) -> int:
    largura, _ = dimensoes(paisagem)
    return int(Emu(largura - 2 * MARGEM_LATERAL).twips)


def _texto_simples(bruto: str) -> str:
    """Texto da celula sem a sintaxe de Markdown, que nao chega a ser impressa."""
    return "".join(f.texto for f in analisar_inline(bruto))


def _largura_em(palavra: str) -> float:
    """Largura estimada da palavra, em "em" (multiplos do corpo)."""
    return sum(LARGURA_EM_MINUSCULA if c.islower() else LARGURA_EM_MAIUSCULA
               for c in palavra)


def _piso_da_coluna(coluna: list[str], util: int) -> int:
    """Largura minima para a palavra mais longa da coluna nao ser partida."""
    mais_larga = 0.0
    for celula in coluna:
        for palavra in _texto_simples(celula).split():
            mais_larga = max(mais_larga, _largura_em(palavra))
    bruto = MARGENS_CELULA_TWIPS + int(mais_larga * TABELA_PT * 20)
    return min(bruto, int(util * TETO_MINIMO_COLUNA))


def escrever_tabela(doc, bloco: Bloco, paisagem: bool = False):
    # a linha de separador e sempre a segunda; nao filtrar por padrao para
    # nao apagar por engano uma linha de dados so com tracos.
    linhas = [l for k, l in enumerate(bloco.linhas) if k != 1]
    matriz = [_celulas(l) for l in linhas]
    n_col = max(len(r) for r in matriz)
    matriz = [r + [""] * (n_col - len(r)) for r in matriz]

    tab = doc.add_table(rows=len(matriz), cols=n_col)
    tab.style = doc.styles["Table Grid"]
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = True

    # larguras preferenciais: proporcionais ao conteudo, amortecidas por sqrt
    # para que uma coluna muito longa nao esmague as restantes. O piso por
    # palavra impede o caso oposto — uma coluna estreita ao lado de uma celula
    # de 1200 caracteres ficava com espaco para menos de uma palavra.
    util = _largura_util(paisagem)
    medidas, pisos = [], []
    for c in range(n_col):
        coluna = [r[c] for r in matriz]
        medidas.append((max(len(x) for x in coluna) or 1) ** 0.5)
        pisos.append(max(int(util * PISO_COLUNA), _piso_da_coluna(coluna, util)))

    soma_pisos = sum(pisos)
    if soma_pisos > util:  # tabela larga demais: encolher os pisos proporcionalmente
        pisos = [int(p * util / soma_pisos) for p in pisos]

    total = sum(medidas)
    larguras = [max(pisos[c], int(util * medidas[c] / total)) for c in range(n_col)]

    # devolver o excesso tirando-o so da folga acima do piso de cada coluna
    excesso = sum(larguras) - util
    if excesso > 0:
        folgas = [larguras[c] - pisos[c] for c in range(n_col)]
        total_folga = sum(folgas)
        if total_folga > 0:
            for c in range(n_col):
                larguras[c] -= min(folgas[c], int(excesso * folgas[c] / total_folga))
        resto = sum(larguras) - util
        if resto > 0:
            maior = larguras.index(max(larguras))
            larguras[maior] -= resto

    for ri, linha in enumerate(matriz):
        for ci, bruto in enumerate(linha):
            celula = tab.cell(ri, ci)
            celula.width = Twips(larguras[ci])
            par = celula.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing = 1.0
            frags = analisar_inline(bruto)
            if ri == 0:
                for f in frags:
                    f.negrito = True
            escrever_fragmentos(par, frags)
            for run in par.runs:
                if run.font.name != FONTE_MONO:
                    run.font.size = Pt(TABELA_PT)
                else:
                    run.font.size = Pt(TABELA_PT - 1)
            if ri == 0:
                _sombrear(celula, FUNDO_CABECALHO_TABELA)
    _repetir_cabecalho(tab.rows[0])

    fecho = doc.add_paragraph()
    fecho.paragraph_format.space_after = Pt(4)
    fecho.paragraph_format.line_spacing = 1.0


def escrever_regua(doc):
    par = doc.add_paragraph()
    pPr = par._p.get_or_add_pPr()
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(8)
    _borda(pPr, "bottom", val="single", sz=6, space=1, color=CINZA_REGUA)


def escrever_citacao(doc, bloco: Bloco):
    filhos = bloco.filhos or [Bloco("paragrafo", texto="")]
    for indice, filho in enumerate(filhos):
        prefixo = ""
        recuo = RECUO_CITACAO
        if filho.tipo == "marca":
            prefixo, recuo = "• ", RECUO_LISTA_CITACAO
        elif filho.tipo == "numero":
            prefixo, recuo = "", RECUO_LISTA_CITACAO
        elif filho.tipo == "titulo":
            pass

        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Twips(recuo)
        par.paragraph_format.space_before = Pt(6 if indice == 0 else 3)
        par.paragraph_format.space_after = Pt(10 if indice == len(filhos) - 1 else 3)
        par.paragraph_format.line_spacing = CORPO_ENTRELINHA
        pPr = par._p.get_or_add_pPr()
        _borda(pPr, "left", val="single", sz=18, space=8, color=AZUL_BARRA_CITACAO)

        frags = analisar_inline(filho.texto)
        if prefixo:
            frags.insert(0, Fragmento(prefixo))
        negrito_titulo = filho.tipo == "titulo"
        for f in frags:
            f.negrito = f.negrito or negrito_titulo
        escrever_fragmentos(par, frags)


def construir(blocos: list[Bloco], paisagem: bool = False) -> "docx.document.Document":
    doc = preparar_documento(paisagem)
    for bloco in blocos:
        if bloco.tipo == "titulo":
            nivel = min(bloco.nivel, 4)
            par = doc.add_heading(level=nivel)
            escrever_fragmentos(par, analisar_inline(bloco.texto))
            for run in par.runs:
                run.bold = True
        elif bloco.tipo == "regua":
            escrever_regua(doc)
        elif bloco.tipo == "tabela":
            escrever_tabela(doc, bloco, paisagem)
        elif bloco.tipo == "citacao":
            escrever_citacao(doc, bloco)
        elif bloco.tipo in ("marca", "numero"):
            base = "List Bullet" if bloco.tipo == "marca" else "List Number"
            estilo = base if bloco.nivel == 1 else f"{base} 2"
            par = doc.add_paragraph(style=estilo)
            escrever_fragmentos(par, analisar_inline(bloco.texto))
        else:
            par = doc.add_paragraph()
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            escrever_fragmentos(par, analisar_inline(bloco.texto))

    titulo = next((b.texto for b in blocos if b.tipo == "titulo" and b.nivel == 1), "")
    doc.core_properties.title = re.sub(r"[`*_]", "", titulo)
    doc.core_properties.author = "LexCV — business/scripts/gerar-docx.py"
    doc.core_properties.comments = (
        "Artefacto derivado. A fonte de verdade e o .md com o mesmo nome. "
        "Nao editar este ficheiro a mao; regenerar com business/scripts/gerar-docx.py."
    )
    return doc


# --------------------------------------------------------------------------
# Escrita determinista
# --------------------------------------------------------------------------

# Um .docx e um zip: as datas das entradas mudam a cada execucao e tornariam
# o binario diferente a cada regeracao, mesmo com .md inalterado. Fixamo-las
# para que "regenerar" seja verificavel — se o .md nao mudou, o .docx nao muda.
DATA_FIXA_ZIP = (1980, 1, 1, 0, 0, 0)


def guardar(doc, destino: Path) -> None:
    import io
    import zipfile

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    with zipfile.ZipFile(buffer) as origem:
        nomes = sorted(origem.namelist())
        conteudos = {n: origem.read(n) for n in nomes}

    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as saida:
        for nome in nomes:
            info = zipfile.ZipInfo(nome, date_time=DATA_FIXA_ZIP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            saida.writestr(info, conteudos[nome])


# --------------------------------------------------------------------------
# Verificacao estrutural
# --------------------------------------------------------------------------

def _todos_os_paragrafos(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for linha in t.rows:
            for celula in linha.cells:
                yield from celula.paragraphs


def verificar(caminho: Path, esperado: dict, paisagem: bool = False) -> list[str]:
    """Reabre o .docx e confere o que e verificavel sem Word/LibreOffice."""
    problemas: list[str] = []
    doc = docx.Document(str(caminho))

    for par in _todos_os_paragrafos(doc):
        if re.search(r"  +", par.text):
            problemas.append(f"espacos multiplos: {par.text[:80]!r}")

    n_tabelas = len(doc.tables)
    if n_tabelas != esperado["tabelas"]:
        problemas.append(f"tabelas: {n_tabelas}, esperadas {esperado['tabelas']}")

    titulos = sum(1 for p in doc.paragraphs if p.style.name.startswith("Heading"))
    if titulos != esperado["titulos"]:
        problemas.append(f"titulos: {titulos}, esperados {esperado['titulos']}")

    listas = sum(1 for p in doc.paragraphs if p.style.name.startswith("List "))
    if listas != esperado["listas"]:
        problemas.append(f"itens de lista: {listas}, esperados {esperado['listas']}")

    for i, t in enumerate(doc.tables):
        if len(t.rows) < 2 or len(t.columns) < 2:
            problemas.append(f"tabela {i} degenerada ({len(t.rows)}x{len(t.columns)})")

    largura, altura = dimensoes(paisagem)
    sec = doc.sections[0]
    # o OOXML guarda a folha em twips, nao em EMU, pelo que a ida e volta perde
    # ate meio twip por dimensao (210 mm = 11905,5 twips). Comparar em twips.
    lidas = (Emu(sec.page_width).twips, Emu(sec.page_height).twips)
    devidas = (Emu(largura).twips, Emu(altura).twips)
    if any(abs(a - b) > 1 for a, b in zip(lidas, devidas)):
        problemas.append(
            f"pagina {lidas[0]}x{lidas[1]} twips, esperada "
            f"A4 {'paisagem' if paisagem else 'retrato'} ({devidas[0]}x{devidas[1]})"
        )
    esperada = WD_ORIENT.LANDSCAPE if paisagem else WD_ORIENT.PORTRAIT
    if sec.orientation != esperada:
        problemas.append(f"orientacao {sec.orientation}, esperada {esperada}")

    # nenhuma coluna pode ficar abaixo do piso: e o defeito que a versao
    # anterior tinha ("Estado" a 0,40 pol, sem espaco para "Previsto").
    piso = int(_largura_util(paisagem) * PISO_COLUNA)
    for i, t in enumerate(doc.tables):
        for c, coluna in enumerate(t.columns):
            largura_col = coluna.cells[0].width
            if largura_col is None:
                continue
            if int(Emu(largura_col).twips) < piso:
                problemas.append(
                    f"tabela {i} coluna {c}: {Emu(largura_col).twips / 1440:.2f} pol, "
                    f"abaixo do piso de {piso / 1440:.2f} pol"
                )

    return problemas


def contar_esperado(blocos: list[Bloco]) -> dict:
    return {
        "tabelas": sum(1 for b in blocos if b.tipo == "tabela"),
        "titulos": sum(1 for b in blocos if b.tipo == "titulo"),
        "listas": sum(1 for b in blocos if b.tipo in ("marca", "numero")),
    }


# --------------------------------------------------------------------------

def main() -> int:
    raiz = Path(__file__).resolve().parent.parent / "documentacao"

    ap = argparse.ArgumentParser(description="Markdown -> .docx (business/documentacao)")
    ap.add_argument("ficheiros", nargs="*", help="ficheiros .md (por omissao: todos)")
    ap.add_argument("--verificar", action="store_true",
                    help="nao escreve; verifica os .docx ja existentes")
    args = ap.parse_args()

    if args.ficheiros:
        fontes = [raiz / Path(f).name for f in args.ficheiros]
    else:
        fontes = sorted(raiz.glob("*.md"))

    if not fontes:
        print(f"nenhum .md encontrado em {raiz}", file=sys.stderr)
        return 1

    falhou = False
    for md in fontes:
        if not md.exists():
            print(f"FALTA   {md}", file=sys.stderr)
            falhou = True
            continue

        destino = md.with_suffix(".docx")
        paisagem = md.name in ORIENTACAO_PAISAGEM
        blocos = analisar(md.read_text(encoding="utf-8"))
        esperado = contar_esperado(blocos)

        if not args.verificar:
            guardar(construir(blocos, paisagem), destino)

        if not destino.exists():
            print(f"FALTA   {destino.name}", file=sys.stderr)
            falhou = True
            continue

        problemas = verificar(destino, esperado, paisagem)
        marca = "ERRO  " if problemas else "OK    "
        tamanho = destino.stat().st_size
        print(f"{marca}  {destino.name:<34} {tamanho/1024:6.1f} KB  "
              f"A4 {'paisagem' if paisagem else 'retrato ':<8} "
              f"titulos={esperado['titulos']:>3} listas={esperado['listas']:>3} "
              f"tabelas={esperado['tabelas']:>2}")
        for p in problemas:
            print(f"          - {p}")
            falhou = True

    if falhou:
        print("\nverificacao estrutural FALHOU", file=sys.stderr)
        return 1

    print("\nverificacao estrutural OK — aparencia NAO confirmada "
          "(sem Word/LibreOffice nesta maquina)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
